"""Dashboard workbooks: presentation only, using the authoritative payload."""
from datetime import date, datetime
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


CURRENCY = '"TZS "#,##0.00;[Red]("TZS "#,##0.00)'
DATE = 'dd mmm yyyy'
INK = '1D2939'
BLUE = '465FFF'
PALE = 'ECF3FF'
LINE = Side(style='thin', color='E4E7EC')


def write_cell(sheet, row, column, value):
    """Keep user text literal, including formula-like names and phone numbers."""
    cell = sheet.cell(row, column)
    if isinstance(value, str):
        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)
        cell.value = value
        cell.data_type = 's'
    else:
        cell.value = value
    cell.font = Font(name='Calibri', size=11, color=INK)
    cell.alignment = Alignment(vertical='top', wrap_text=True)
    if isinstance(value, (date, datetime)):
        cell.number_format = DATE
    return cell


def build_dashboard_excel(report):
    workbook = Workbook()
    workbook.properties.creator = report['prepared_by']
    workbook.properties.title = 'MarketFlow · Dashboard Performance Report'
    sheet = workbook.active
    sheet.title = 'Overview'
    sheet.sheet_view.showGridLines = False
    for column, width in zip('ABCDEF', [24, 21, 21, 24, 21, 21]):
        sheet.column_dimensions[column].width = width

    def banner(row, text, size, fill=BLUE, color='FFFFFF'):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = write_cell(sheet, row, 1, text)
        cell.font = Font(name='Calibri', size=size, bold=True, color=color)
        cell.fill = PatternFill('solid', fgColor=fill)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        sheet.row_dimensions[row].height = 36 if size > 15 else 28

    banner(1, 'MarketFlow', 24)
    banner(2, 'Dashboard Performance Report', 16)
    scope, metrics = report['scope'], report['metrics']
    identity = [
        ('Marketer', report['marketer']), ('Commission period', scope.period.name),
        ('Start date', scope.period.start), ('End date', scope.period.end),
        ('Location', report['location']),
        ('Generated at (' + report['generated_at'].tzname() + ')', report['generated_at'].replace(tzinfo=None)),
        ('Prepared by', report['prepared_by']),
    ]
    for row, (label, value) in enumerate(identity, 4):
        write_cell(sheet, row, 1, label).font = Font(name='Calibri', bold=True, color=INK)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        cell = write_cell(sheet, row, 2, value)
        if isinstance(value, datetime):
            cell.number_format = 'dd mmm yyyy hh:mm'
        sheet.row_dimensions[row].height = max(32, (len(str(value)) // 100 + str(value).count('\n') + 1) * 16)
    banner(12, 'Dashboard summary · Amounts in TZS', 14, PALE, INK)
    kpis = [('Total Sales', 'sales'), ('Total Expenditure', 'expenditure'),
            ('Base Amount', 'base'), ('Amount After Base', 'excess'),
            ('Commission Rate', 'rate'), ('Total Commission', 'commission'),
            ('Total Customers', 'customers')]
    for row, (label, key) in enumerate(kpis, 14):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        sheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
        left = write_cell(sheet, row, 1, label)
        left.font = Font(name='Calibri', bold=True, color=INK)
        value = metrics.get(key)
        if value is None:
            value = 'Commission policy not configured' if key in ('base', 'excess', 'rate', 'commission') and metrics['sales'] is not None else 'Not permitted'
        elif key == 'rate':
            value = value / 100
        right = write_cell(sheet, row, 4, value)
        right.number_format = '0.00%' if key == 'rate' else '#,##0' if key == 'customers' else CURRENCY
        right.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
        right.font = Font(name='Calibri', size=14, bold=True, color=BLUE)
        for cell in (left, right):
            cell.fill = PatternFill('solid', fgColor=PALE if row % 2 == 0 else 'F9FAFB')
            cell.border = Border(bottom=LINE)
        sheet.row_dimensions[row].height = 36
    notes = [report['explanation']]
    if metrics['configured']:
        notes.append(f"max(TZS {metrics['sales']:,.2f} − TZS {metrics['base']:,.2f}, 0) × {metrics['rate']:g}% = TZS {metrics['commission']:,.2f}")
    else:
        notes.append('Commission policy not configured' if metrics['sales'] is not None else 'Sales and commission: not permitted')
    notes += ['Expenditure does not reduce commission. Sales and commission cover all eligible period sales across locations.',
              'The selected location and its permitted descendants filter expenditure and customers only.',
              'MarketFlow · Confidential · Amounts in TZS']
    for row, note in enumerate(notes, 23):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        write_cell(sheet, row, 1, note)
        sheet.row_dimensions[row].height = 32
    sheet.freeze_panes = 'A14'
    sheet.print_options.horizontalCentered = True
    sheet.print_area = 'A1:F27'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    widths = {'Sales': [38, 36, 25, 18], 'Expenditure': [36, 25, 25, 30, 18, 23],
              'Customers': [36, 25, 32, 18]}
    for section in report['sections']:
        sheet = workbook.create_sheet(section['title'])
        sheet.sheet_view.showGridLines = False
        for col, (label, width) in enumerate(zip(section['headers'], widths[section['title']]), 1):
            cell = write_cell(sheet, 1, col, label)
            cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor=BLUE)
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.row_dimensions[1].height = 30
        for row, values in enumerate(section['rows'], 2):
            for col, value in enumerate(values, 1):
                cell = write_cell(sheet, row, col, value)
                if section['headers'][col - 1] == 'Amount':
                    cell.number_format = CURRENCY
                    cell.alignment = Alignment(horizontal='right', vertical='top')
                elif section['headers'][col - 1] == 'Phone':
                    cell.number_format = '@'
                if row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='F9FAFB')
            # Accommodate wrapped names without clipping in Excel's fixed-height rows.
            lines = max((len(str(v)) // max(12, int(w) - 3) + str(v).count('\n') + 1
                         for v, w in zip(values, widths[section['title']])), default=1)
            sheet.row_dimensions[row].height = max(30, lines * 16)
        last_column = get_column_letter(len(section['headers']))
        last_row = len(section['rows']) + 1
        if section['rows']:
            table = Table(displayName='Dashboard' + section['title'], ref=f'A1:{last_column}{last_row}')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            sheet.add_table(table)
        else:
            write_cell(sheet, 2, 1, 'No records in this scope.' if section['available'] else 'Not permitted')
            sheet.row_dimensions[2].height = 30
        sheet.auto_filter.ref = f'A1:{last_column}{last_row}'
        sheet.freeze_panes = 'A2'
        sheet.print_title_rows = '1:1'
        sheet.print_area = f'A1:{last_column}{max(2, last_row)}'
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    for sheet in workbook:
        sheet.oddFooter.left.text = 'MarketFlow · Confidential · Amounts in TZS'
        sheet.oddFooter.right.text = 'Page &P of &N'
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
