import uuid
from datetime import date, timedelta
from decimal import Decimal
from importlib import import_module
from unittest.mock import patch

from django.apps import apps
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.db import connection, IntegrityError, transaction
from django.db.migrations.executor import MigrationExecutor
from django.http import QueryDict
from django.test import TestCase, TransactionTestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone

from .dashboard import get_dashboard_scope, get_dashboard_metrics
from .models import CommissionPeriod, CommissionPolicy, Sale, Customer, Expenditure, ExpenditureType, Location
from .services import get_default_commission_period, calculate_commission


class DashboardTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_superuser('dashboard-admin')
        group = Group.objects.create(name='Marketer')
        group.permissions.set(Permission.objects.filter(content_type__app_label='api', codename__in=[
            'view_reports', 'view_sale', 'view_customer', 'view_expenditure', 'add_customer']))
        cls.a = get_user_model().objects.create_user('alice', first_name='Alice')
        cls.b = get_user_model().objects.create_user('bob', first_name='Bob')
        cls.a.groups.add(group)
        cls.b.groups.add(group)
        cls.period = CommissionPeriod.objects.create(name='September 2026', start=date(2026,9,1), end=date(2026,9,30), is_default=True)
        cls.other = CommissionPeriod.objects.create(name='August 2026', start=date(2026,8,1), end=date(2026,8,31))
        cls.policy = CommissionPolicy.objects.create(marketer=cls.a, period=cls.period, base=3000000, rate=3)
        cls.root = Location.objects.create(name='Dar es Salaam', normalized_name='dar', level='region', identity='dar')
        cls.child = Location.objects.create(name='Kinondoni', normalized_name='kinondoni', level='district', parent=cls.root, identity='kin')
        cls.private = Location.objects.create(name='Private', normalized_name='private', level='region', creator=cls.b, identity='private')
        cls.type = ExpenditureType.objects.create(name='Travel', identity='travel')

    def owned(self, user=None, **extra):
        return dict(marketer=user or self.a, created_by=self.admin, updated_by=self.admin, date=date(2026,9,15), **extra)

    def sale(self, amount=7000000, status='confirmed', user=None, **extra):
        return Sale.objects.create(reference=str(uuid.uuid4()), amount=amount, status=status, **self.owned(user, **extra))

    def expense(self, amount=850000, status='recorded', location=None, user=None):
        return Expenditure.objects.create(title='Travel cost', expenditure_type=self.type, amount=amount, status=status, location=location, **self.owned(user))

    def customer(self, status='finalized', location=None, user=None):
        return Customer.objects.create(name='Customer', phone='+255712345678', comment='Interested', status=status, location=location, **self.owned(user))

    def dashboard(self, user=None, **params):
        self.client.force_login(user or self.a)
        return self.client.get('/dashboard/', params)

    def details(self, metric, **params):
        return self.client.get(reverse('api:dashboard-detail', args=[metric]), params, HTTP_HX_REQUEST='true')

    def test_default_switch_and_database_constraint(self):
        self.other.is_default=True
        self.other.save()
        self.period.refresh_from_db()
        self.assertFalse(self.period.is_default)
        self.assertEqual(get_default_commission_period(), self.other)
        with self.assertRaises(IntegrityError), transaction.atomic():
            CommissionPeriod.objects.filter(pk=self.period.pk).update(is_default=True)
        self.assertEqual(CommissionPeriod.objects.filter(is_default=True).count(), 1)

    def test_settings_switch_and_permissions(self):
        self.client.force_login(self.admin)
        response=self.client.post(reverse('api:edit', args=['periods',self.other.pk]), {
            'name':self.other.name,'start':self.other.start,'end':self.other.end,'is_default':'on'})
        self.assertEqual(response.status_code,302)
        self.assertEqual(get_default_commission_period(),self.other)
        self.client.force_login(self.a)
        self.assertEqual(self.client.post(reverse('api:edit',args=['periods',self.period.pk]),{}).status_code,403)

    def test_default_period_cannot_be_overridden(self):
        self.sale()
        old=self.sale(99);old.date=date(2026,8,15);old.save()
        for value in [self.other.pk, 'bad', '99999']:
            r=self.dashboard(period=value)
            self.assertEqual(r.status_code,200)
            self.assertEqual(r.context['period'],self.period)
            self.assertEqual(r.context['metrics']['sales'],7000000)
            detail=self.details('sales',period=value)
            self.assertEqual(detail.context['metrics']['sales'],7000000)
            self.assertNotIn(old,list(detail.context['page']))

    def test_missing_default_is_setup_not_all_time(self):
        CommissionPeriod.objects.update(is_default=False)
        self.sale()
        r=self.dashboard()
        self.assertContains(r,'Commission period is not configured.')
        self.assertContains(r,'Contact an administrator')
        self.assertNotContains(r,'7,000,000')
        r=self.dashboard(self.admin)
        self.assertContains(r,'Set commission period →')
        self.assertEqual(r.context['setup_url'],reverse('api:edit',args=['periods',self.period.pk]))

    def test_automatic_selection_and_invalid_marketers(self):
        self.assertEqual(self.dashboard().context['scope'].marketer,self.a)
        self.assertEqual(self.dashboard(self.admin).context['scope'].marketer,self.a)
        self.assertEqual(self.dashboard(self.admin,marketer=self.b.pk).context['scope'].marketer,self.b)
        for value in ['bad','999999999999999999999999999',self.b.pk]:
            self.assertEqual(self.dashboard(marketer=value).status_code,400)
            self.assertEqual(self.details('customers',marketer=value).status_code,400)
        self.assertEqual(self.dashboard(self.admin,marketer='bad').status_code,400)
        self.assertEqual(self.dashboard(marketer=[self.a.pk,self.b.pk]).status_code,400)

    def test_current_cohort_member_preferred(self):
        self.a.user_permissions.add(*Permission.objects.filter(content_type__app_label='api',codename__in=['view_all_marketing','view_all_reports']))
        self.assertEqual(self.dashboard(self.a).context['scope'].marketer,self.a)
        self.b.is_active=False;self.b.save()
        self.assertEqual(self.dashboard(self.admin).context['scope'].marketer,self.a)

    def test_exact_formula_status_and_expenditure_independence(self):
        self.sale(); self.sale(9000000,'draft'); self.sale(9000000,'void'); self.sale(8000000,user=self.b)
        self.expense(); self.expense(123,'void')
        self.customer(); self.customer('archived')
        metrics=self.dashboard().context['metrics']
        self.assertEqual((metrics['sales'],metrics['excess'],metrics['commission'],metrics['rate']),(7000000,4000000,120000,3))
        self.assertEqual(metrics['expenditure'],850000)
        self.assertEqual(metrics['customers'],1)
        self.expense(999999)
        self.assertEqual(self.dashboard().context['metrics']['commission'],120000)

    def test_empty_scope_and_bounded_location_lookup(self):
        r=self.dashboard()
        self.assertEqual((r.context['metrics']['sales'],r.context['metrics']['commission'],r.context['metrics']['customers']),(0,0,0))
        for metric in ['sales','expenditure','customers']:
            self.assertContains(self.details(metric),'in this scope.')
        for i in range(35):
            Location.objects.create(name='Region '+str(i),normalized_name='region'+str(i),level='region',identity='region'+str(i))
        data=self.client.get('/lookups/location/').json()
        self.assertEqual(len(data['results']),30)
        self.assertTrue(data['more'])
        self.assertNotIn(self.private.pk,[row['id'] for row in data['results']])

    def test_below_base_zero_base_and_rounding(self):
        sale=self.sale(2000000)
        m=self.dashboard().context['metrics']
        self.assertEqual((m['excess'],m['commission']),(0,0))
        self.policy.base=0;self.policy.save()
        sale.amount=Decimal('1.50');sale.save()
        m=self.dashboard().context['metrics']
        self.assertEqual((m['excess'],m['commission']),(Decimal('1.50'),Decimal('.05')))
        self.assertEqual(calculate_commission(Decimal('0'),Decimal('0'),Decimal('3'))['commission'],0)

    def test_location_descendants_and_official_commission(self):
        inside=self.customer(location=self.child)
        self.customer();self.customer(user=self.b,location=self.child)
        self.expense(location=self.child);self.expense(12)
        self.sale(customer=inside);self.sale(100)
        r=self.dashboard(location=self.root.pk)
        m=r.context['metrics']
        self.assertEqual((m['customers'],m['expenditure'],m['sales'],m['commission']),(1,850000,7000100,Decimal('120003')))
        self.assertContains(r,'Sales and commission cover all locations')
        self.assertEqual(self.details('sales',location=self.root.pk).context['page'].paginator.count,2)
        self.assertEqual(self.details('customers',location=self.root.pk).context['page'].paginator.count,1)
        self.assertEqual(self.details('expenditure',location=self.root.pk).context['metrics']['expenditure'],850000)
        for value in [self.private.pk,'bad','999999']:
            self.assertEqual(self.dashboard(location=value).status_code,400)
            self.assertEqual(self.details('customers',location=value).status_code,400)

    def test_expenditure_list_location_display_and_filter(self):
        self.expense(location=self.child);self.expense(12)
        self.client.force_login(self.a)
        r=self.client.get('/marketing/expenditures/',{'location':self.root.pk})
        self.assertContains(r,'Kinondoni')
        self.assertContains(r,'data-lookup="location"')
        self.assertIn('location',r.context['columns'])
        self.assertEqual(r.context['amount_total'],850000)

    def test_popups_match_cards_and_paginate(self):
        for _ in range(23):
            self.sale(100);self.expense(10);self.customer()
        self.sale(1,'draft');self.expense(1,'void');self.customer('archived')
        card=self.dashboard().context['metrics']
        for metric in ['sales','expenditure','customers']:
            first=self.details(metric)
            second=self.details(metric,page=2)
            rows=list(first.context['page'])+list(second.context['page'])
            self.assertEqual(len(rows),23)
            self.assertEqual(len(first.context['page']),20)
            self.assertEqual(len({r.pk for r in rows}),23)
            actual=len(rows) if metric=='customers' else sum(r.amount for r in rows)
            self.assertEqual(actual,card[metric])
            self.assertEqual(first.context['metrics'][metric],card[metric])
            self.assertNotContains(first,'<!doctype html>')
            self.assertContains(first,'data-modal')
            direct=self.client.get(reverse('api:dashboard-detail',args=[metric]))
            self.assertContains(direct,'<!doctype html>')
        for metric in ['commission','after-base']:
            r=self.details(metric)
            self.assertContains(r,'Commissionable amount cannot be below zero.')
            self.assertNotContains(r,'<!doctype html>')

    def test_detail_allowlist_auth_and_permissions(self):
        self.assertEqual(self.client.get('/dashboard/details/customers/').status_code,302)
        self.dashboard()
        self.assertEqual(self.details('users').status_code,404)
        self.assertEqual(self.details('commission',marketer=self.b.pk).status_code,400)
        self.a.groups.clear()
        self.assertEqual(self.details('sales').status_code,403)

    def test_missing_or_inactive_policy_quick_setup(self):
        self.sale(user=self.b)
        r=self.dashboard(self.b)
        self.assertContains(r,'Not configured',count=2)
        self.assertNotContains(r,'0%')
        self.assertContains(self.details('commission'),'administrator configuration')
        r=self.dashboard(self.admin,marketer=self.b.pk)
        self.assertIn('marketer='+str(self.b.pk),r.context['policy_url'])
        form=self.client.get(r.context['policy_url'],HTTP_HX_REQUEST='true')
        self.assertEqual(form.context['form'].initial['marketer'],self.b)
        self.assertEqual(form.context['form'].initial['period'],self.period)
        self.policy.active=False;self.policy.save()
        r=self.dashboard(self.admin,marketer=self.a.pk)
        self.assertEqual(r.context['policy_url'],reverse('api:edit',args=['commissions',self.policy.pk]))

    def test_dashboard_html_only_two_filters_and_semantic_cards(self):
        r=self.dashboard(self.admin)
        html=r.content.decode()
        from html.parser import HTMLParser
        class Elements(HTMLParser):
            def __init__(self):super().__init__();self.selects=[];self.cards=[]
            def handle_starttag(self,tag,attrs):
                attrs=dict(attrs)
                if tag=='select':self.selects.append(attrs)
                if 'dashboard-metric' in attrs.get('class','').split():self.cards.append((tag,attrs))
        elements=Elements();elements.feed(html)
        self.assertEqual([s['name'] for s in elements.selects],['marketer','location'])
        self.assertEqual(len(elements.cards),6)
        for tag,attrs in elements.cards:
            self.assertEqual(tag,'a');self.assertIn('href',attrs);self.assertIn('aria-label',attrs);self.assertIn('data-modal',attrs)
        for text in ['finance-region','data-period','effective rate','Marketer performance','All permitted marketers','multiple']:
            self.assertNotIn(text,html)
        fragment=self.client.get('/dashboard/',HTTP_HX_REQUEST='true')
        self.assertNotContains(fragment,'<!doctype html>')
        restored=self.client.get('/dashboard/',HTTP_HX_REQUEST='true',HTTP_HX_HISTORY_RESTORE_REQUEST='true')
        self.assertContains(restored,'<!doctype html>')

    def test_non_sales_reader_retains_permitted_cards(self):
        user=get_user_model().objects.create_user('reader')
        user.user_permissions.set(Permission.objects.filter(content_type__app_label='api',codename__in=['view_reports','view_customer']))
        r=self.dashboard(user)
        self.assertContains(r,'Total Customers')
        self.assertNotContains(r,'Total Sales')
        self.assertEqual(self.details('sales').status_code,403)

    def test_bounded_queries_and_lookup(self):
        self.admin.get_all_permissions()
        q=QueryDict('marketer='+str(self.a.pk))
        with CaptureQueriesContext(connection) as captured:
            get_dashboard_metrics(get_dashboard_scope(self.admin,q))
        initial=len(captured)
        for i in range(40):
            user=get_user_model().objects.create_user('extra'+str(i));user.groups.add(Group.objects.get(name='Marketer'))
        with CaptureQueriesContext(connection) as captured:
            scope=get_dashboard_scope(self.admin,q);get_dashboard_metrics(scope)
        self.assertLessEqual(len(captured),initial)
        self.assertLessEqual(len(scope.options),31)
        self.dashboard(self.admin)
        result=self.client.get('/lookups/dashboard-marketer/',{'q':'Bob','period':'bad'}).json()
        self.assertEqual([row['id'] for row in result['results']],[self.b.pk])
        self.dashboard()
        self.assertEqual(self.client.get('/lookups/dashboard-marketer/',{'q':'Bob'}).json()['results'],[])

    def allow_export(self, user=None):
        user = user or self.a
        user.user_permissions.add(Permission.objects.get(content_type__app_label='api', codename='export_reports'))
        self.client.force_login(user)

    def export_response(self, format, **params):
        return self.client.get(reverse('api:dashboard-export-' + format), params)

    def pdf_text(self, response):
        from io import BytesIO
        from pypdf import PdfReader
        self.assertEqual(response.status_code, 200, response.content[:500])
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF-'))
        reader = PdfReader(BytesIO(response.content))
        return '\n'.join(page.extract_text() for page in reader.pages)

    def workbook(self, response):
        from io import BytesIO
        from openpyxl import load_workbook
        self.assertEqual(response.status_code, 200, response.content[:500])
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.sheetnames, ['Overview', 'Sales', 'Expenditure', 'Customers'])
        return book

    def test_dashboard_vendor_absent_but_record_and_crud_preserved(self):
        record = self.expense(location=self.child)
        record.vendor = 'Unique private vendor'
        record.payment_method = 'Mobile Money'
        record.save()
        original = Expenditure.objects.filter(pk=record.pk).values().get()
        self.client.force_login(self.admin)
        response = self.details('expenditure', marketer=self.a.pk)
        for text in ['Title', 'Type', 'Amount', 'Location', 'Date', 'Payment method', 'Travel cost', 'Mobile Money']:
            self.assertContains(response, text)
        self.assertNotContains(response, 'Vendor')
        self.assertNotContains(response, record.vendor)
        self.assertEqual(response.context['page'].paginator.count, 1)
        self.assertEqual(Expenditure.objects.filter(pk=record.pk).values().get(), original)
        detail = self.client.get(reverse('api:detail', args=['expenditures', record.pk]))
        self.assertContains(detail, record.vendor)
        edit = self.client.get(reverse('api:edit', args=['expenditures', record.pk]))
        self.assertContains(edit, 'name="vendor"')
        self.assertContains(edit, record.vendor)
        self.assertContains(self.details('expenditure', marketer=self.b.pk), 'colspan="6"')
        self.assertContains(self.details('sales', marketer=self.b.pk), 'colspan="4"')

    def test_export_card_permission_and_all_endpoint_guards(self):
        self.assertNotContains(self.dashboard(), 'Export Report')
        routes = ['dashboard-export', 'dashboard-export-pdf', 'dashboard-export-excel']
        for route in routes:
            self.assertEqual(self.client.get(reverse('api:' + route)).status_code, 403)
        self.allow_export()
        self.assertContains(self.client.get('/dashboard/'), 'Export Report')
        self.a.groups.clear()
        for route in routes:
            self.assertEqual(self.client.get(reverse('api:' + route)).status_code, 403)
        self.client.logout()
        for route in routes:
            self.assertEqual(self.client.get(reverse('api:' + route)).status_code, 302)
            self.assertEqual(self.client.get(reverse('api:' + route), HTTP_HX_REQUEST='true').status_code, 401)

    def test_export_chooser_validated_scope_and_full_page_fallback(self):
        self.client.force_login(self.admin)
        params = dict(marketer=self.a.pk, location=self.child.pk, period=self.other.pk)
        url = reverse('api:dashboard-export')
        fragment = self.client.get(url, params, HTTP_HX_REQUEST='true')
        self.assertEqual(fragment.context['scope'].marketer, self.a)
        self.assertEqual(fragment.context['scope'].location, self.child)
        for text in ['Get PDF', 'Get Excel', 'Alice', 'September 2026', '01 Sep 2026', '30 Sep 2026', 'Location: Kinondoni']:
            self.assertContains(fragment, text)
        for text in ['<!doctype', '<select', 'period=', 'August 2026', 'data-modal']:
            self.assertNotContains(fragment, text)
        for format in ['pdf', 'excel']:
            self.assertContains(fragment, reverse('api:dashboard-export-' + format) + f'?marketer={self.a.pk}&amp;location={self.child.pk}')
        self.assertContains(self.client.get(url, params), '<!doctype html>')
        self.assertContains(self.client.get(url, params, HTTP_HX_REQUEST='true', HTTP_HX_HISTORY_RESTORE_REQUEST='true'), '<!doctype html>')
        self.assertContains(self.client.get(url), 'Location: All permitted locations')

    def test_export_scope_security_and_method(self):
        self.allow_export()
        for route in ['dashboard-export', 'dashboard-export-pdf', 'dashboard-export-excel']:
            url = reverse('api:' + route)
            for params in [dict(marketer=self.b.pk), dict(marketer='bad'), dict(marketer=[self.a.pk, self.b.pk]),
                           dict(location=self.private.pk), dict(location='bad'), dict(location=[self.child.pk, self.root.pk])]:
                self.assertEqual(self.client.get(url, params).status_code, 400)
            self.assertEqual(self.client.post(url).status_code, 405)

    def test_exports_exact_dashboard_numbers_period_and_location(self):
        self.allow_export()
        customer = self.customer(location=self.child)
        self.customer(); self.customer('archived', location=self.child)
        self.customer(user=self.b, location=self.child)
        self.sale(customer=customer)
        self.sale(9000000, 'draft'); self.sale(9000000, 'void'); self.sale(9000000, user=self.b)
        old = self.sale(42); old.date = self.other.start; old.save()
        params = dict(marketer=self.a.pk, location=self.root.pk, period=self.other.pk)
        before = self.workbook(self.export_response('excel', **params))
        self.assertEqual(before['Overview']['D19'].value, 120000)
        record = self.expense(location=self.child)
        record.vendor = 'Export secret vendor'; record.payment_method = 'Cash'; record.save()
        self.expense(12); self.expense(55, 'void', location=self.child); self.expense(44, user=self.b, location=self.child)
        pdf = self.export_response('pdf', **params)
        text = self.pdf_text(pdf)
        for value in ['MarketFlow', 'Dashboard Performance Report', 'Alice', 'September 2026', '01 Sep 2026',
                      '30 Sep 2026', '7,000,000.00', '850,000.00', '3,000,000.00', '4,000,000.00',
                      '3.00%', '120,000.00', '1 customer', 'Travel cost', 'Kinondoni', 'Cash', 'Page 1']:
            self.assertIn(value, text)
        for value in ['Vendor', record.vendor, 'August 2026', str(old.reference)]:
            self.assertNotIn(value, text)
        self.assertEqual(pdf['Content-Disposition'], 'attachment; filename="marketflow-dashboard-alice-september-2026.pdf"')
        excel = self.export_response('excel', **params)
        self.assertEqual(excel['Content-Disposition'], 'attachment; filename="marketflow-dashboard-alice-september-2026.xlsx"')
        book = self.workbook(excel)
        overview = book['Overview']
        self.assertEqual([overview[f'D{r}'].value for r in range(14, 21)], [7000000, 850000, 3000000, 4000000, .03, 120000, 1])
        self.assertEqual(overview['B4'].value, 'Alice')
        self.assertEqual(overview['B5'].value, 'September 2026')
        self.assertEqual(overview['B6'].value.date(), self.period.start)
        self.assertEqual(overview['B7'].value.date(), self.period.end)
        self.assertEqual(overview['B8'].value, 'Dar es Salaam')
        self.assertEqual(overview['B10'].value, 'Alice')
        expense = book['Expenditure']
        self.assertEqual(list(expense.values)[0], ('Title', 'Type', 'Amount', 'Location', 'Date', 'Payment Method'))
        self.assertEqual(expense.max_row, 2)
        self.assertEqual(expense['C2'].value, 850000)
        self.assertEqual(expense['C2'].data_type, 'n')
        self.assertIn('TZS', expense['C2'].number_format)
        self.assertEqual(expense['E2'].value.date(), record.date)
        self.assertEqual(book['Sales'].max_row, 2)
        self.assertEqual(book['Customers'].max_row, 2)
        self.assertEqual(book['Customers']['B2'].value, customer.phone)
        self.assertEqual(book['Customers']['B2'].data_type, 's')
        for sheet in list(book)[1:]:
            self.assertEqual(sheet.freeze_panes, 'A2')
            self.assertTrue(sheet.auto_filter.ref)
            self.assertEqual(sheet.print_title_rows, '$1:$1')
        record.refresh_from_db()
        self.assertEqual(record.vendor, 'Export secret vendor')

    def test_exports_missing_policy_and_default_do_not_imply_zero(self):
        self.allow_export()
        self.policy.active = False; self.policy.save()
        self.assertIn('Commission policy not configured', self.pdf_text(self.export_response('pdf')))
        book = self.workbook(self.export_response('excel'))
        for row in [16, 17, 18, 19]:
            self.assertEqual(book['Overview'][f'D{row}'].value, 'Commission policy not configured')
        CommissionPeriod.objects.update(is_default=False)
        for format in ['pdf', 'excel']:
            self.assertContains(self.export_response(format), 'Configure a default commission period', status_code=400)

    def test_export_location_never_reduces_sales_entitlement(self):
        self.allow_export()
        inside = self.customer(location=self.child)
        outside = self.customer()
        self.sale(2000000, customer=inside)
        self.sale(5000000, customer=outside)
        self.expense(850000, location=self.child)
        self.expense(9000000)
        book = self.workbook(self.export_response('excel', location=self.root.pk))
        self.assertEqual(book['Sales'].max_row, 3)
        self.assertEqual(book['Customers'].max_row, 2)
        self.assertEqual(book['Expenditure'].max_row, 2)
        self.assertEqual([book['Overview'][f'D{r}'].value for r in [14, 15, 17, 19]],
                         [7000000, 850000, 4000000, 120000])
        text = self.pdf_text(self.export_response('pdf', location=self.root.pk))
        for amount in ['7,000,000.00', '2,000,000.00', '5,000,000.00', '4,000,000.00', '120,000.00', '850,000.00']:
            self.assertIn(amount, text)
        self.assertNotIn('9,000,000.00', text)

    def test_exports_respect_individual_model_permissions(self):
        reader = get_user_model().objects.create_user('export-reader')
        reader.user_permissions.set(Permission.objects.filter(content_type__app_label='api',
            codename__in=['view_reports', 'export_reports', 'view_customer']))
        self.customer(user=reader)
        self.sale(user=reader)
        self.client.force_login(reader)
        text = self.pdf_text(self.export_response('pdf'))
        self.assertIn('Not permitted', text)
        self.assertNotIn('7,000,000.00', text)
        book = self.workbook(self.export_response('excel'))
        self.assertEqual(book['Sales']['A2'].value, 'Not permitted')
        self.assertEqual(book['Overview']['D14'].value, 'Not permitted')
        self.assertEqual(book['Customers']['A2'].value, 'Customer')

    def test_export_limit_combines_sections_and_never_renders_partial_files(self):
        self.allow_export()
        self.sale(); self.expense(); self.customer()
        with patch('config.api.dashboard.DASHBOARD_EXPORT_LIMIT', 2), \
             patch('config.api.pdf.build_dashboard_pdf') as pdf, patch('config.api.excel.build_dashboard_excel') as excel:
            for format in ['pdf', 'excel']:
                self.assertContains(self.export_response(format), 'No partial report was generated', status_code=422)
            pdf.assert_not_called(); excel.assert_not_called()
        with patch('config.api.dashboard.DASHBOARD_EXPORT_LIMIT', 3):
            self.workbook(self.export_response('excel'))

    def test_export_text_is_literal_and_filename_is_safe(self):
        self.allow_export()
        self.a.first_name = '=HYPERLINK("https://example.com") / Alice'; self.a.save()
        customer = self.customer()
        customer.name = '=1+1'; customer.save()
        self.sale(customer=customer)
        record = self.expense(); record.title = '@SUM(1,1)'; record.save()
        response = self.export_response('excel')
        book = self.workbook(response)
        for cell in [book['Overview']['B4'], book['Customers']['A2'], book['Sales']['B2'], book['Expenditure']['A2']]:
            self.assertEqual(cell.data_type, 's')
        self.assertNotIn('/', response['Content-Disposition'])
        self.assertNotIn('=', response['Content-Disposition'].split('filename=')[1])

    def test_export_rows_not_paginated_and_queries_not_per_record(self):
        from .dashboard import get_dashboard_export
        self.allow_export()
        self.sale(); self.expense(location=self.child); self.customer(location=self.child)
        scope = get_dashboard_scope(self.a, QueryDict(''))
        self.a.get_all_permissions()
        with CaptureQueriesContext(connection) as queries:
            get_dashboard_export(scope)
        initial = len(queries)
        for _ in range(25):
            c = self.customer(location=self.child); self.sale(customer=c); self.expense(location=self.child)
        with CaptureQueriesContext(connection) as queries:
            payload = get_dashboard_export(scope)
        self.assertLessEqual(len(queries), initial)
        self.assertEqual([len(s['rows']) for s in payload['sections']], [26, 26, 26])
        book = self.workbook(self.export_response('excel'))
        for sheet in list(book)[1:]:
            self.assertEqual(sheet.max_row, 27)
        text = self.pdf_text(self.export_response('pdf'))
        self.assertIn('26 sales', text)
        self.assertIn('26 expenditure records', text)
        self.assertIn('26 customers', text)
        self.assertIn('Page 2', text)


class DefaultPeriodMigrationTests(TransactionTestCase):
    def test_existing_installation_priority_and_idempotency(self):
        old=('api','0004_alter_commissionpolicy_period_and_more')
        new=('api','0005_commissionperiod_is_default_and_more')
        executor=MigrationExecutor(connection)
        executor.migrate([old])
        try:
            old_apps=executor.loader.project_state([old]).apps
            Period=old_apps.get_model('api','CommissionPeriod')
            today=timezone.localdate()
            current=Period.objects.create(name='Current',start=today,end=today,closed=True)
            future=Period.objects.create(name='Future open',start=today+timedelta(days=40),end=today+timedelta(days=50))
            executor=MigrationExecutor(connection);executor.migrate([new])
            self.assertEqual(get_default_commission_period().pk,current.pk)
            migrate=import_module('config.api.migrations.0005_commissionperiod_is_default_and_more').choose_default
            with connection.schema_editor() as editor:
                migrate(apps,editor)
                self.assertEqual(CommissionPeriod.objects.count(),2)
                CommissionPeriod.objects.update(is_default=False)
                with patch('django.utils.timezone.localdate',return_value=today-timedelta(days=1)):
                    migrate(apps,editor)
                    self.assertEqual(get_default_commission_period().pk,future.pk)
                    CommissionPeriod.objects.update(is_default=False,closed=True)
                    migrate(apps,editor)
                    self.assertEqual(get_default_commission_period().pk,future.pk)
        finally:
            MigrationExecutor(connection).migrate([new])
